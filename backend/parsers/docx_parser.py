"""DOCX parser → paragraphs + table cells as text."""
from __future__ import annotations

import io

from parsers.base import ParsedDocument, infer_doc_type


def parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []

    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    full = "\n\n".join(parts)
    return ParsedDocument(
        text=full,
        doc_type=infer_doc_type(full),
        page_count=1,
        metadata={"parser": "python-docx", "tables": len(doc.tables)},
    )
